import streamlit as st
import zipfile, io, re, os, tempfile
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

st.set_page_config(page_title="Tender Comparative Statement", page_icon="📑", layout="wide")
st.title("Tender Technical Bid Comparative Statement")
st.caption("ZIP files upload karein → documents se details extract karein → review karein → Word comparative download karein.")

st.info("Privacy: Yeh app aapke computer par locally chalane ke liye banaya gaya hai. Files kisi server par bhejne ki zaroorat nahi, jab aap ise apne PC par run karte hain.")

uploaded = st.file_uploader(
    "Agency-wise Technical Bid ZIP files select karein",
    type=["zip"],
    accept_multiple_files=True,
    help="Har ZIP ek agency ki honi chahiye. File name se agency ka naam prefill hoga."
)

FIELDS = [
    ("Tender Fee / DD", ["DD number / amount / date"]),
    ("EMD / FDR", ["EMD/FDR number, amount, date"]),
    ("PAN Card", ["PAN number"]),
    ("GST", ["GSTIN"]),
    ("Electrical Contractor Licence", ["Licence number, issue/renewal/expiry"]),
    ("Bank Solvency", ["Amount and certificate date"]),
    ("Agency Registration", ["Class/category, date and validity"]),
]

def safe_name(name):
    stem = Path(name).stem
    stem = re.sub(r'[_-]+', ' ', stem)
    stem = re.sub(r'\bTECHNICAL BID\b', '', stem, flags=re.I)
    stem = re.sub(r'\b\d{5,}\b', '', stem)
    return re.sub(r'\s+', ' ', stem).strip() or "Agency"

def extract_pdf(data):
    text = ""
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
    except Exception:
        pass
    # Optional OCR for scanned PDFs if Tesseract + pdf2image are installed.
    if len(re.sub(r'\s', '', text)) < 30:
        try:
            import pytesseract
            from pdf2image import convert_from_bytes
            images = convert_from_bytes(data, dpi=200)
            text = "\n".join(pytesseract.image_to_string(im, lang="eng") for im in images)
        except Exception:
            pass
    return text

def extract_docx(data):
    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.replace("\n", " ") for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""

def extract_text_from_zip(upload):
    result = []
    try:
        zf = zipfile.ZipFile(upload)
    except Exception:
        return result, "ZIP file open nahi hui."
    for info in zf.infolist():
        if info.is_dir():
            continue
        name = info.filename
        ext = Path(name).suffix.lower()
        try:
            data = zf.read(info)
        except Exception:
            continue
        txt = ""
        if ext == ".pdf":
            txt = extract_pdf(data)
        elif ext == ".docx":
            txt = extract_docx(data)
        elif ext in (".txt", ".csv"):
            txt = data.decode("utf-8", errors="ignore")
        elif ext == ".doc":
            txt = ""
        if txt.strip():
            result.append((name, txt))
    return result, None

def first_match(patterns, text, flags=re.I):
    for pattern in patterns:
        m = re.search(pattern, text, flags)
        if m:
            return re.sub(r'\s+', ' ', m.group(0)).strip()
    return ""

def parse_details(documents):
    alltext = "\n".join(f"\n--- FILE: {n} ---\n{t}" for n,t in documents)
    flat = re.sub(r'[ \t]+', ' ', alltext)
    # GSTIN/PAN patterns
    gst = first_match([r'\b\d{2}[A-Z]{5}\d{4}[A-Z][A-Z0-9]Z[A-Z0-9]\b'], flat)
    pan = first_match([r'\b[A-Z]{5}\d{4}[A-Z]\b'], flat)
    # Find DD-like references near DD / demand draft labels
    dd = ""
    m = re.search(r'(?:D\.?\s*D\.?|DEMAND DRAFT|DD NO\.?|DD NUMBER)[^\n]{0,100}', flat, re.I)
    if m: dd = m.group(0).strip()
    # Common EMD/FDR context
    emd = ""
    m = re.search(r'(?:E\.?\s*M\.?\s*D\.?|FDR|FIXED DEPOSIT RECEIPT|BANK GUARANTEE)[^\n]{0,130}', flat, re.I)
    if m: emd = m.group(0).strip()
    # Licence context
    licence = ""
    m = re.search(r'(?:ELECTRICAL CONTRACTOR|CONTRACTOR LICEN[CS]E|LICEN[CS]E NO\.?)[^\n]{0,180}', flat, re.I)
    if m: licence = m.group(0).strip()
    # Solvency context
    solvency = ""
    m = re.search(r'(?:SOLVENCY|SOLVENT)[^\n]{0,180}', flat, re.I)
    if m: solvency = m.group(0).strip()
    # Registration context
    reg = ""
    m = re.search(r'(?:R\s*&\s*B|REGISTRATION|REGISTERED AS|CLASS\s*[A-Z])[^\n]{0,180}', flat, re.I)
    if m: reg = m.group(0).strip()

    return {
        "Tender Fee / DD": dd or "Manual verification required",
        "EMD / FDR": emd or "Manual verification required",
        "PAN Card": pan or "Not detected — verify scan",
        "GST": gst or "Not detected — verify scan",
        "Electrical Contractor Licence": licence or "Manual verification required",
        "Bank Solvency": solvency or "Manual verification required",
        "Agency Registration": reg or "Manual verification required",
    }, alltext

def build_docx(rows, tender_id, include_review_notes=True):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = Cm(1.2); sec.bottom_margin = Cm(1.2)
    sec.left_margin = Cm(1.2); sec.right_margin = Cm(1.2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TECHNICAL BID – COMPARATIVE STATEMENT")
    r.bold = True; r.font.size = Pt(14)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run(f"Tender ID: {tender_id or '________________'}").bold = True

    headers = ["Sr. No.", "Agency Name"] + [f[0] for f in FIELDS]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(1.0), Cm(3.2)] + [Cm(3.1)]*len(FIELDS)
    for i,h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True; run.font.size = Pt(8)
    for idx,row in enumerate(rows,1):
        cells = table.add_row().cells
        vals = [str(idx), row["agency"]] + [row["details"].get(f[0],"") for f in FIELDS]
        for i,val in enumerate(vals):
            cells[i].text = str(val)
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                for run in p.runs: run.font.size = Pt(7)
    doc.add_paragraph("")
    note = doc.add_paragraph()
    run = note.add_run("Note: This statement is machine-assisted. Please verify all extracted details against original documents before official use.")
    run.italic = True; run.font.size = Pt(8)
    doc.add_paragraph("")
    sig = doc.add_table(rows=1, cols=3)
    sig.style = "Table Grid"
    for c, txt in zip(sig.rows[0].cells, ["Senior Clerk", "Divisional Accounts Officer", "Executive Engineer"]):
        c.text = "\n\n\n" + txt
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.font.size = Pt(9)
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out

if uploaded:
    st.subheader("Agency details")
    tender_id = st.text_input("Tender ID / Tender No.", value="335762")
    rows = []
    for i, file in enumerate(uploaded):
        agency_default = safe_name(file.name)
        with st.expander(f"{i+1}. {file.name}", expanded=True):
            agency = st.text_input("Agency name", value=agency_default, key=f"agency_{i}")
            docs, err = extract_text_from_zip(file)
            if err:
                st.error(err)
            st.write(f"Documents with readable text: **{len(docs)}**")
            if docs:
                st.caption("Files processed: " + ", ".join(Path(n).name for n,_ in docs[:12]))
            else:
                st.warning("Readable text nahi mila. ZIP mein scans/image-only PDFs ho sakte hain. OCR setup ke baad dobara try karein ya details manually fill karein.")
            details, alltext = parse_details(docs)
            st.caption("Extracted values draft hain — original documents se verify karna zaroori hai.")
            edited = {}
            cols = st.columns(2)
            for j,(field, _) in enumerate(FIELDS):
                with cols[j % 2]:
                    edited[field] = st.text_area(field, value=details[field], height=72, key=f"field_{i}_{j}")
            rows.append({"agency": agency, "details": edited, "raw_text": alltext})
    st.divider()
    if st.button("Generate Word Comparative Statement", type="primary"):
        out = build_docx(rows, tender_id)
        st.download_button(
            "Download Comparative Statement (.docx)",
            data=out,
            file_name=f"Comparative_Statement_{re.sub(r'[^A-Za-z0-9_-]','_',tender_id or 'Tender')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        st.success("Word file ready. Download button par click karein.")
else:
    st.markdown("### Shuru kaise karein")
    st.markdown("1. Neeche diye gaye software package ko apne Windows computer par extract karein.\n2. Python install karein (Python 3.10 ya newer).\n3. `install.bat` ek baar chalayein.\n4. `run.bat` chalayein.\n5. Browser mein khule page par agency ZIP files upload karein.")
    st.warning("Scanned PDFs ke liye OCR ki zaroorat ho sakti hai. OCR ke bina bhi text-based PDFs aur DOCX files process hongi. Automatic extraction 100% accurate nahi hota; final approval se pehle original documents verify karein.")
